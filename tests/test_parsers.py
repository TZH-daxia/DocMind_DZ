"""多格式解析器。"""

from pathlib import Path

import pytest

from app.collector.parsers import ParserError, build_default_registry
from tests.conftest import MASTER_PAGE, write_pdf


def test_pdf_parser_keeps_text_and_images(settings, tmp_path: Path):
    source = write_pdf(tmp_path / "master.pdf", [MASTER_PAGE])
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    parsed = build_default_registry().get(".pdf").parse(source, out_dir, settings)

    assert parsed.parser == "pdf"
    assert len(parsed.pages) == 1
    assert "MASTER AIR WAYBILL" in parsed.pages[0].text
    assert parsed.pages[0].text_source == "text_layer"
    assert parsed.pages[0].image_path is not None
    assert parsed.pages[0].image_path.stat().st_size > 0


def test_docx_parser_extracts_paragraphs_and_tables(settings, tmp_path: Path):
    from docx import Document

    document = Document()
    document.add_paragraph("HOUSE AIR WAYBILL")
    document.add_paragraph("HAWB NO.: HAWB0001")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Origin"
    table.cell(0, 1).text = "PVG"
    table.cell(1, 0).text = "Destination"
    table.cell(1, 1).text = "LAX"
    source = tmp_path / "house.docx"
    document.save(source)

    out_dir = tmp_path / "out"
    out_dir.mkdir()
    parsed = build_default_registry().get(".docx").parse(source, out_dir, settings)

    text = parsed.pages[0].text
    assert "HAWB NO.: HAWB0001" in text
    assert "Origin | PVG" in text
    assert parsed.pages[0].text_source == "docx"


def test_xlsx_parser_flattens_sheets(settings, tmp_path: Path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MANIFEST"
    sheet.append(["HAWB NO.", "ORIGIN", "DEST"])
    sheet.append(["HAWB0001", "PVG", "LAX"])
    source = tmp_path / "manifest.xlsx"
    workbook.save(source)

    out_dir = tmp_path / "out"
    out_dir.mkdir()
    parsed = build_default_registry().get(".xlsx").parse(source, out_dir, settings)

    assert "Sheet: MANIFEST" in parsed.pages[0].text
    assert "HAWB0001 | PVG | LAX" in parsed.pages[0].text


def test_image_parser_creates_page(settings, tmp_path: Path):
    from PIL import Image

    source = tmp_path / "scan.png"
    Image.new("RGB", (300, 200), "white").save(source)
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    parsed = build_default_registry().get(".png").parse(source, out_dir, settings)

    assert len(parsed.pages) == 1
    assert parsed.pages[0].text == ""
    assert parsed.pages[0].image_path is not None
    assert parsed.pages[0].image_path.exists()


def test_unsupported_extension_raises():
    with pytest.raises(ParserError) as excinfo:
        build_default_registry().get(".txt")
    assert excinfo.value.code == "FILE_TYPE_NOT_SUPPORTED"
