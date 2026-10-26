"""端到端管线（注入假网关，不调用真实模型）。"""

import asyncio

import pytest

from app.schemas.file import UploadedDocument
from app.service.task_service import TaskError, TaskService
from tests.conftest import (
    MANIFEST_PAGE,
    MASTER_PAGE,
    house_page,
    make_service,
    write_pdf,
)


async def _run_task(service: TaskService, task_id: str) -> None:
    service.start(task_id)
    for _ in range(200):
        snapshot = service.get_snapshot(task_id)
        if snapshot.status in {"succeeded", "failed"}:
            return
        await asyncio.sleep(0.05)
    raise AssertionError("任务未在预期时间内结束")


async def test_pipeline_skips_master_and_compares_manifest(settings, tmp_path):
    pdf = write_pdf(
        tmp_path / "bundle.pdf",
        [
            MASTER_PAGE,
            house_page("HAWB0001"),
            house_page("HAWB0002"),
            MANIFEST_PAGE,
        ],
    )
    service = make_service(settings)
    snapshot = service.create_task(
        UploadedDocument(filename="bundle.pdf", content=pdf.read_bytes())
    )
    await _run_task(service, snapshot.task_id)

    snapshot = service.get_snapshot(snapshot.task_id)
    assert snapshot.status == "succeeded", snapshot.error
    result = service.get_result(snapshot.task_id)

    assert result["page_count"] == 4
    assert len(result["master_bills"]) == 1
    master = result["master_bills"][0]
    assert master["skipped"] is True
    assert master["page_indexes"] == [1]
    assert master["master_no"] == "999-12345678"

    house_bills = result["house_bills"]
    assert len(house_bills) == 2
    assert [bill["page_indexes"] for bill in house_bills] == [[2], [3]]
    # 主单页没有被当成任何一份分单的输入
    assert all(1 not in bill["page_indexes"] for bill in house_bills)

    assert result["manifest"] is not None
    assert len(result["manifest"]["entries"]) == 3

    summary = result["summary"]
    assert summary["total_house_bills"] == 2
    assert summary["total_manifest_entries"] == 3
    assert summary["consistent"] == 1
    assert summary["inconsistent"] == 1
    assert summary["manifest_only"] == 1

    inconsistent = next(
        item for item in result["comparisons"] if item["status"] == "inconsistent"
    )
    weight = next(field for field in inconsistent["fields"] if field["key"] == "weight")
    assert weight["status"] == "mismatch"


async def test_pipeline_succeeds_with_master_only_file(settings, tmp_path):
    pdf = write_pdf(tmp_path / "master_only.pdf", [MASTER_PAGE, MANIFEST_PAGE])
    service = make_service(settings)
    snapshot = service.create_task(
        UploadedDocument(filename="master_only.pdf", content=pdf.read_bytes())
    )
    await _run_task(service, snapshot.task_id)

    snapshot = service.get_snapshot(snapshot.task_id)
    assert snapshot.status == "succeeded"
    result = service.get_result(snapshot.task_id)
    assert result["house_bills"] == []
    assert len(result["master_bills"]) == 1
    assert any("主单" in warning for warning in result["warnings"])


async def test_manifest_only_file_succeeds_with_hint(settings, tmp_path):
    """只有舱单时同样跳过抽取，并提示补充分单与主单。"""

    pdf = write_pdf(tmp_path / "manifest_only.pdf", [MANIFEST_PAGE])
    service = make_service(settings)
    snapshot = service.create_task(
        UploadedDocument(filename="manifest_only.pdf", content=pdf.read_bytes())
    )
    await _run_task(service, snapshot.task_id)

    assert service.get_snapshot(snapshot.task_id).status == "succeeded"
    result = service.get_result(snapshot.task_id)
    assert result["house_bills"] == []
    assert any("只包含舱单" in warning for warning in result["warnings"])


async def test_master_only_file_skips_extract_and_verify_nodes(settings, tmp_path):
    """只有主单时，图应直接跳到结果装配，不执行抽取与核对节点。"""

    pdf = write_pdf(tmp_path / "master_single.pdf", [MASTER_PAGE])
    service = make_service(settings)
    snapshot = service.create_task(
        UploadedDocument(filename="master_single.pdf", content=pdf.read_bytes())
    )
    await _run_task(service, snapshot.task_id)

    assert service.get_snapshot(snapshot.task_id).status == "succeeded"
    events = service.read_events(snapshot.task_id)
    skipped = {event.step for event in events if event.phase == "skipped"}
    started = {event.step for event in events if event.phase == "started"}
    assert skipped == {"抽取分单与舱单字段", "核对分单与舱单"}
    assert "抽取分单与舱单字段" not in started
    assert "核对分单与舱单" not in started


async def test_pipeline_accepts_house_pdf_with_manifest_xlsx(settings, tmp_path):
    from openpyxl import Workbook

    pdf = write_pdf(
        tmp_path / "house.pdf", [house_page("HAWB0001"), house_page("HAWB0002")]
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MANIFEST"
    sheet.append(["CARGO MANIFEST"])
    sheet.append(["HAWB NO.", "ORIGIN", "DEST", "PCS"])
    sheet.append(["HAWB0001", "PVG", "LAX", "12"])
    xlsx = tmp_path / "manifest.xlsx"
    workbook.save(xlsx)

    service = make_service(settings)
    snapshot = service.create_task(
        [
            UploadedDocument(filename="house.pdf", content=pdf.read_bytes()),
            UploadedDocument(filename="manifest.xlsx", content=xlsx.read_bytes()),
        ]
    )
    await _run_task(service, snapshot.task_id)

    snapshot = service.get_snapshot(snapshot.task_id)
    assert snapshot.status == "succeeded", snapshot.error
    result = service.get_result(snapshot.task_id)
    assert result["page_count"] == 3
    assert [page["source_file"] for page in result["pages"]] == [
        "house.pdf",
        "house.pdf",
        "manifest.xlsx",
    ]
    assert len(result["house_bills"]) == 2
    assert result["manifest"] is not None
    assert result["manifest"]["entries"]


def test_create_task_rejects_empty_upload_list(settings):
    service = make_service(settings)
    with pytest.raises(TaskError) as excinfo:
        service.create_task([])
    assert excinfo.value.code == "FILE_EMPTY"


def test_reject_unsupported_extension(settings):
    service = make_service(settings)
    with pytest.raises(TaskError) as excinfo:
        service.create_task(UploadedDocument(filename="a.txt", content=b"hello"))
    assert excinfo.value.code == "FILE_TYPE_NOT_SUPPORTED"


def test_reject_fake_pdf(settings):
    service = make_service(settings)
    with pytest.raises(TaskError) as excinfo:
        service.create_task(UploadedDocument(filename="a.pdf", content=b"not a pdf"))
    assert excinfo.value.code == "FILE_CONTENT_INVALID"


def test_reject_too_large(settings):
    service = make_service(settings.model_copy(update={"max_file_size_mb": 1}))
    with pytest.raises(TaskError) as excinfo:
        service.create_task(
            UploadedDocument(filename="big.pdf", content=b"%PDF" + b"0" * (2 * 1024 * 1024))
        )
    assert excinfo.value.code == "FILE_TOO_LARGE"
