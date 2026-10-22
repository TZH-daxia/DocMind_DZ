"""多格式解析器：把上传文件统一成「页面列表（文本 + 可选图片）」。

设计要点：
- 所有格式都产出 `ParsedDocument`，下游管线只认页面，不再关心文件后缀；
- PDF 同时保留文本层与页面图片：有可靠文本层时直接用文本（快且准），
  扫描件/文本层缺失时由管线调用 VLM 转写；
- Word/Excel 优先用 LibreOffice 转 PDF，再复用 PDF 解析器（有页图、文本层一致）；
  LibreOffice 不可用时回退 python-docx / openpyxl / xlrd 的纯文本解析；
- 图片文件（png/jpg 等）暂不支持，上传时按不支持类型拒绝。
"""

from __future__ import annotations

import logging
import shutil
import subprocess
from dataclasses import dataclass, field
from pathlib import Path
from typing import Protocol

from app.config import Settings

logger = logging.getLogger(__name__)

# Word / Excel 无分页概念，按行数切成"逻辑页"，保证单页内容量可控
DOCX_LINES_PER_PAGE = 55
SHEET_ROWS_PER_PAGE = 80


class ParserError(RuntimeError):
    """解析失败，message 为面向前端的中文原因。"""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class ParsedPage:
    index: int
    text: str = ""
    text_source: str = "empty"
    image_path: Path | None = None


@dataclass
class ParsedDocument:
    pages: list[ParsedPage] = field(default_factory=list)
    parser: str = ""
    warnings: list[str] = field(default_factory=list)

    @property
    def page_count(self) -> int:
        return len(self.pages)


class DocumentParser(Protocol):
    name: str
    extensions: tuple[str, ...]

    def parse(self, source: Path, out_dir: Path, settings: Settings) -> ParsedDocument: ...


def _normalize_text(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    cleaned: list[str] = []
    for line in lines:
        if not line.strip() and cleaned and not cleaned[-1].strip():
            continue
        cleaned.append(line)
    return "\n".join(cleaned).strip()


def _paginate_lines(lines: list[str], per_page: int) -> list[str]:
    pages: list[str] = []
    for start in range(0, max(len(lines), 1), per_page):
        chunk = _normalize_text("\n".join(lines[start : start + per_page]))
        if chunk:
            pages.append(chunk)
    return pages or [""]


class PdfParser:
    name = "pdf"
    extensions = (".pdf",)

    def parse(self, source: Path, out_dir: Path, settings: Settings) -> ParsedDocument:
        try:
            import pymupdf
        except ImportError as exc:  # pragma: no cover - 依赖缺失
            raise ParserError("PARSER_UNAVAILABLE", "缺少 PyMuPDF 依赖，无法解析 PDF") from exc

        document = ParsedDocument(parser="pdf")
        zoom = settings.render_dpi / 72.0
        matrix = pymupdf.Matrix(zoom, zoom)
        try:
            with pymupdf.open(source) as pdf:
                total = pdf.page_count
                if total > settings.max_pages:
                    document.warnings.append(
                        f"文件共 {total} 页，仅解析前 {settings.max_pages} 页"
                    )
                for offset in range(min(total, settings.max_pages)):
                    page = pdf.load_page(offset)
                    text = _normalize_text(page.get_text("text"))
                    image_path = out_dir / f"page_{offset + 1:03d}.png"
                    try:
                        pixmap = page.get_pixmap(matrix=matrix, alpha=False)
                        pixmap.save(image_path)
                    except Exception as exc:  # noqa: BLE001 - 单页渲染失败不阻断
                        logger.warning("第 %s 页渲染失败：%s", offset + 1, exc)
                        image_path = None
                    document.pages.append(
                        ParsedPage(
                            index=offset + 1,
                            text=text,
                            text_source="text_layer" if text else "empty",
                            image_path=image_path,
                        )
                    )
        except ParserError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise ParserError("PARSE_FAILED", f"PDF 解析失败：{exc}") from exc
        return document


class DocxParser:
    name = "docx"
    extensions = (".docx",)

    def parse(self, source: Path, out_dir: Path, settings: Settings) -> ParsedDocument:
        try:
            from docx import Document as WordDocument
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError as exc:  # pragma: no cover
            raise ParserError("PARSER_UNAVAILABLE", "缺少 python-docx 依赖") from exc

        document = ParsedDocument(parser="docx")
        try:
            word = WordDocument(str(source))
            lines: list[str] = []
            for child in word.element.body.iterchildren():
                tag = child.tag.rsplit("}", 1)[-1]
                if tag == "p":
                    text = Paragraph(child, word).text.strip()
                    if text:
                        lines.append(text)
                elif tag == "tbl":
                    table = Table(child, word)
                    for row in table.rows:
                        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        if any(cells):
                            lines.append(" | ".join(cells))
            for index, text in enumerate(_paginate_lines(lines, DOCX_LINES_PER_PAGE), start=1):
                document.pages.append(
                    ParsedPage(index=index, text=text, text_source="docx")
                )
            if not document.pages:
                document.pages.append(ParsedPage(index=1, text="", text_source="empty"))
        except ParserError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise ParserError("PARSE_FAILED", f"Word 解析失败：{exc}") from exc
        return document


class ExcelParser:
    name = "xlsx"
    extensions = (".xlsx",)

    def parse(self, source: Path, out_dir: Path, settings: Settings) -> ParsedDocument:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise ParserError("PARSER_UNAVAILABLE", "缺少 openpyxl 依赖") from exc

        document = ParsedDocument(parser="xlsx")
        try:
            workbook = load_workbook(source, read_only=True, data_only=True)
            page_index = 0
            for sheet in workbook.worksheets:
                lines = [f"# Sheet: {sheet.title}"]
                for row in sheet.iter_rows(values_only=True):
                    cells = ["" if value is None else str(value).strip() for value in row]
                    while cells and not cells[-1]:
                        cells.pop()
                    if any(cells):
                        lines.append(" | ".join(cells))
                for chunk in _paginate_lines(lines, SHEET_ROWS_PER_PAGE):
                    page_index += 1
                    document.pages.append(
                        ParsedPage(index=page_index, text=chunk, text_source="xlsx")
                    )
            workbook.close()
            if not document.pages:
                document.pages.append(ParsedPage(index=1, text="", text_source="empty"))
        except Exception as exc:  # noqa: BLE001
            raise ParserError("PARSE_FAILED", f"Excel 解析失败：{exc}") from exc
        return document


class XlsParser:
    name = "xls"
    extensions = (".xls",)

    def parse(self, source: Path, out_dir: Path, settings: Settings) -> ParsedDocument:
        try:
            import xlrd
        except ImportError as exc:  # pragma: no cover
            raise ParserError("PARSER_UNAVAILABLE", "缺少 xlrd 依赖") from exc

        document = ParsedDocument(parser="xls")
        try:
            book = xlrd.open_workbook(str(source))
            page_index = 0
            for sheet in book.sheets():
                lines = [f"# Sheet: {sheet.name}"]
                for row_index in range(sheet.nrows):
                    cells = [str(value).strip() for value in sheet.row_values(row_index)]
                    while cells and not cells[-1]:
                        cells.pop()
                    if any(cells):
                        lines.append(" | ".join(cells))
                for chunk in _paginate_lines(lines, SHEET_ROWS_PER_PAGE):
                    page_index += 1
                    document.pages.append(
                        ParsedPage(index=page_index, text=chunk, text_source="xlsx")
                    )
            if not document.pages:
                document.pages.append(ParsedPage(index=1, text="", text_source="empty"))
        except Exception as exc:  # noqa: BLE001
            raise ParserError("PARSE_FAILED", f"Excel(xls) 解析失败：{exc}") from exc
        return document


def _libreoffice_binary() -> str | None:
    """定位 LibreOffice 可执行文件（PATH 优先，其次常见安装位置）。"""

    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for candidate in (
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
        Path("/usr/bin/soffice"),
    ):
        if candidate.exists():
            return str(candidate)
    return None


def _office_to_pdf(source: Path, out_dir: Path) -> Path | None:
    """用 LibreOffice headless 把 Office 文档转成 PDF；失败返回 None 由调用方回退。"""

    binary = _libreoffice_binary()
    if binary is None:
        return None
    command = [
        binary,
        "--headless",
        "--norestore",
        "--convert-to",
        "pdf",
        "--outdir",
        str(out_dir),
        str(source),
    ]
    try:
        subprocess.run(command, check=True, capture_output=True, timeout=180)
    except (subprocess.SubprocessError, OSError) as exc:
        logger.warning("Office 转 PDF 失败：%s（%s）", source.name, exc)
        return None
    target = out_dir / f"{source.stem}.pdf"
    return target if target.exists() else None


class OfficePdfParser:
    """Office 文档（doc/docx/xls/xlsx）先转 PDF，再复用 PDF 解析器。

    这样 Word/Excel 也有页面图片与文本层，左侧预览与 PDF 走同一条链路；
    LibreOffice 不可用或转换失败时回退到纯文本解析（无页图）。
    """

    name = "office_pdf"
    extensions = (".doc", ".docx", ".xls", ".xlsx")

    def parse(self, source: Path, out_dir: Path, settings: Settings) -> ParsedDocument:
        suffix = source.suffix.lower()
        if settings.office_render_enabled:
            pdf_path = _office_to_pdf(source, out_dir)
            if pdf_path is not None:
                document = PdfParser().parse(pdf_path, out_dir, settings)
                document.parser = f"{suffix.lstrip('.')}_pdf"
                return document

        if suffix == ".doc":
            raise ParserError(
                "LEGACY_FORMAT_UNSUPPORTED",
                "旧版 .doc 需要安装 LibreOffice 才能解析，请另存为 .docx 后重试",
            )
        fallback: DocumentParser
        if suffix == ".docx":
            fallback = DocxParser()
        elif suffix == ".xlsx":
            fallback = ExcelParser()
        else:
            fallback = XlsParser()
        document = fallback.parse(source, out_dir, settings)
        document.warnings.append("未启用或未安装 LibreOffice，已按纯文本解析，无页面图片")
        return document


class ParserRegistry:
    """后缀 → 解析器注册表（新增格式只需 register）。"""

    def __init__(self) -> None:
        self._parsers: dict[str, DocumentParser] = {}

    def register(self, parser: DocumentParser) -> None:
        for extension in parser.extensions:
            self._parsers[extension.lower()] = parser

    def get(self, suffix: str) -> DocumentParser:
        parser = self._parsers.get(suffix.lower())
        if parser is None:
            raise ParserError("FILE_TYPE_NOT_SUPPORTED", f"暂不支持的文件类型：{suffix}")
        return parser

    def supports(self, suffix: str) -> bool:
        return suffix.lower() in self._parsers

    @property
    def extensions(self) -> tuple[str, ...]:
        return tuple(sorted(self._parsers))


def build_default_registry() -> ParserRegistry:
    registry = ParserRegistry()
    # Office 文档统一「转 PDF → 复用 PDF 解析器」，保证有页面图片；失败时解析器内部回退
    registry.register(OfficePdfParser())
    registry.register(PdfParser())
    return registry
